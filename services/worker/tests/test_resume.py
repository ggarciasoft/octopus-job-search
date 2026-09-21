"""Resume assembly, validation and rendering (M3, PR07).

The assertions that matter are the refusals. A generated CV is a document
someone will send to an employer under their own name, so the interesting
cases are the ones where the honest output is *less* than the model offered:
a dropped number, a dropped employer, a project that stayed a project.

AT10 -- "Generated CV adding a numeric claim is blocked or flagged" -- is
`test_a_number_absent_from_the_facts_is_dropped_and_flagged`.
"""

from __future__ import annotations

import io
from typing import Any, Literal

import pytest
from docx import Document as DocxDocument

from job_getter_worker.contracts.generated import (
    JobRequirement,
    ProfileFact,
    ResumeBullet,
    ResumeContact,
    ResumeDocument,
    ResumeEntry,
    ResumeSection,
)
from job_getter_worker.resume.document import MissingContactError, build_document
from job_getter_worker.resume.docx_render import render_docx
from job_getter_worker.resume.pdf_render import build_html
from job_getter_worker.resume.validation import validate_document

CONTACT_ID = "11111111-1111-4111-8111-111111111111"
EXPERIENCE_ID = "22222222-2222-4222-8222-222222222222"
EXPERIENCE_ID_2 = "22222222-2222-4222-8222-222222222223"
SKILL_ID = "33333333-3333-4333-8333-333333333333"
PROJECT_ID = "44444444-4444-4444-8444-444444444444"
CERT_ID = "55555555-5555-4555-8555-555555555555"
EDUCATION_ID = "66666666-6666-4666-8666-666666666666"
LANGUAGE_ID = "77777777-7777-4777-8777-777777777777"
UNKNOWN_ID = "99999999-9999-4999-8999-999999999999"

NOW = "2026-09-21T00:00:00.000Z"


def fact(fact_id: str, kind: str, value: dict[str, Any], *, confirmed: bool = True) -> ProfileFact:
    return ProfileFact.model_validate(
        {
            "id": fact_id,
            "kind": kind,
            "value": value,
            "source_file_id": None,
            "source_excerpt": None,
            "confirmed": confirmed,
            "revision": 1,
            "supersedes_id": None,
            "created_at": NOW,
            "updated_at": NOW,
        }
    )


def contact_fact() -> ProfileFact:
    return fact(
        CONTACT_ID,
        "contact",
        {
            "full_name": "Ada Lovelace",
            "email": "ada@example.invalid",
            "phone": "+1 555 0100",
            "city": "Montevideo",
            "country": "Uruguay",
            "links": [{"label": "Site", "url": "https://ada.example.invalid"}],
        },
    )


def experience_fact(
    fact_id: str = EXPERIENCE_ID,
    employer: str = "Orbital Foods",
    title: str = "Senior Backend Engineer",
    bullets: list[str] | None = None,
) -> ProfileFact:
    return fact(
        fact_id,
        "experience",
        {
            "employer": employer,
            "title": title,
            "start_month": "2019-03",
            "end_month": None,
            "current": True,
            "employment_type": "full_time",
            "bullets": [{"text": text, "evidence_reference": "p1"} for text in (bullets or [])],
            "skills": [],
        },
    )


def skill_fact(name: str = "Python", fact_id: str = SKILL_ID) -> ProfileFact:
    return fact(
        fact_id,
        "skill",
        {"canonical_name": name, "aliases": [], "user_declared_proficiency": None, "years": None},
    )


def baseline_facts() -> list[ProfileFact]:
    return [contact_fact(), experience_fact(), skill_fact()]


def entry(**overrides: Any) -> ResumeEntry:
    base: dict[str, Any] = {
        "title": "Senior Backend Engineer",
        "organization": "Orbital Foods",
        "start_month": "2019-03",
        "end_month": None,
        "current": True,
        "detail": None,
        "bullets": [],
        "fact_ids": [EXPERIENCE_ID],
    }
    base.update(overrides)
    return ResumeEntry.model_validate(base)


def document_with(
    section_kind: Literal[
        "summary", "skills", "experience", "projects", "education", "certifications", "languages"
    ],
    entries: list[ResumeEntry],
) -> ResumeDocument:
    return ResumeDocument(
        schema_version=1,
        language="en",
        contact=ResumeContact(
            full_name="Ada Lovelace",
            email="ada@example.invalid",
            phone=None,
            location=None,
            links=[],
            fact_ids=[CONTACT_ID],
        ),
        sections=[ResumeSection(kind=section_kind, heading="Experience", entries=entries)],
    )


# ---------------------------------------------------------------------------
# The deterministic document
# ---------------------------------------------------------------------------


def test_the_document_invents_nothing_without_a_provider() -> None:
    """Every string is copied from a fact or is a fixed heading."""
    document = build_document(baseline_facts(), "en")

    assert document.contact.full_name == "Ada Lovelace"
    kinds = [section.kind for section in document.sections]
    assert kinds == ["skills", "experience"]
    experience = next(s for s in document.sections if s.kind == "experience")
    assert experience.entries[0].organization == "Orbital Foods"


def test_a_profile_with_no_contact_fact_refuses_rather_than_going_out_unnamed() -> None:
    """A CV with no name is not a CV, and a placeholder name is an invention.

    This is the case every other fixture in this file accidentally avoids by
    always including a contact fact. It reached a live stack as a retryable
    INTERNAL_ERROR before anything here caught it.
    """
    with pytest.raises(MissingContactError):
        build_document([skill_fact(), experience_fact()], "en")


def test_a_malformed_contact_fact_is_treated_as_absent() -> None:
    broken = fact(CONTACT_ID, "contact", {"full_name": "", "email": "not-an-email"})

    with pytest.raises(MissingContactError):
        build_document([broken, skill_fact()], "en")


def test_empty_sections_are_omitted_not_left_as_bare_headings() -> None:
    document = build_document([contact_fact()], "en")

    assert document.contact.full_name == "Ada Lovelace"
    assert document.sections == []


def test_headings_follow_the_requested_language() -> None:
    english = build_document(baseline_facts(), "en")
    spanish = build_document(baseline_facts(), "es")

    assert [s.heading for s in english.sections] == ["Skills", "Experience"]
    assert [s.heading for s in spanish.sections] == ["Habilidades", "Experiencia"]


def test_a_job_reorders_skills_without_adding_or_removing_any() -> None:
    facts = [
        contact_fact(),
        skill_fact("Python", SKILL_ID),
        skill_fact("Kubernetes", "33333333-3333-4333-8333-333333333334"),
    ]
    requirements = [
        JobRequirement(text="Kubernetes in production", kind="required", evidence_excerpt="k8s")
    ]

    plain = build_document(facts, "en")
    tailored = build_document(facts, "en", requirements)

    plain_skills = [
        e.title or "" for e in next(s for s in plain.sections if s.kind == "skills").entries
    ]
    tailored_skills = [
        e.title or "" for e in next(s for s in tailored.sections if s.kind == "skills").entries
    ]
    assert sorted(plain_skills) == sorted(tailored_skills)
    assert tailored_skills[0] == "Kubernetes"


def test_a_draft_fact_never_reaches_the_document() -> None:
    draft = skill_fact("Rust", "33333333-3333-4333-8333-333333333335")
    draft = draft.model_copy(update={"confirmed": False})

    document = build_document([contact_fact(), draft], "en")

    assert document.sections == []


def test_a_project_is_never_given_an_employer() -> None:
    project = fact(
        PROJECT_ID,
        "project",
        {
            "name": "Analytical Engine",
            "role": "Author",
            "url": None,
            "start_month": None,
            "end_month": None,
            "bullets": [],
        },
    )
    document = build_document([contact_fact(), project], "en")
    projects = next(s for s in document.sections if s.kind == "projects")

    assert projects.entries[0].organization is None


# ---------------------------------------------------------------------------
# Validation: the refusals
# ---------------------------------------------------------------------------


def test_a_number_absent_from_the_facts_is_dropped_and_flagged() -> None:
    """AT10. A model adding "by 40%" to a fact that never said 40."""
    facts = [contact_fact(), experience_fact(bullets=["Ran the ingestion pipeline"])]
    document = document_with(
        "experience",
        [
            entry(
                bullets=[
                    ResumeBullet(
                        text="Cut ingestion latency by 40% across the pipeline",
                        fact_ids=[EXPERIENCE_ID],
                    )
                ]
            )
        ],
    )

    outcome = validate_document(document, facts)

    codes = [finding.code for finding in outcome.findings]
    assert "NUMBER_NOT_IN_FACTS" in codes
    assert outcome.document.sections[0].entries[0].bullets == []
    assert not outcome.passed
    flagged = next(f for f in outcome.findings if f.code == "NUMBER_NOT_IN_FACTS")
    assert flagged.removed is True
    assert "40%" in (flagged.excerpt or "")


def test_a_number_the_facts_do_contain_survives() -> None:
    facts = [
        contact_fact(),
        experience_fact(bullets=["Ran the ingestion pipeline for 12 teams"]),
    ]
    document = document_with(
        "experience",
        [entry(bullets=[ResumeBullet(text="Supported 12 teams", fact_ids=[EXPERIENCE_ID])])],
    )

    outcome = validate_document(document, facts)

    assert outcome.passed
    assert outcome.document.sections[0].entries[0].bullets[0].text == "Supported 12 teams"


def test_an_employer_the_profile_never_mentions_is_dropped() -> None:
    document = document_with("experience", [entry(organization="Initech")])

    outcome = validate_document(document, baseline_facts())

    assert "NAME_NOT_IN_FACTS" in [f.code for f in outcome.findings]
    assert outcome.document.sections[0].entries[0].organization is None
    assert not outcome.passed


def test_a_date_the_profile_never_stated_is_dropped() -> None:
    document = document_with("experience", [entry(start_month="2010-01")])

    outcome = validate_document(document, baseline_facts())

    assert "DATE_NOT_IN_FACTS" in [f.code for f in outcome.findings]
    assert outcome.document.sections[0].entries[0].start_month is None


def test_two_roles_merged_into_one_entry_are_refused() -> None:
    facts = [
        contact_fact(),
        experience_fact(),
        experience_fact(EXPERIENCE_ID_2, employer="Acme", title="Staff Engineer"),
    ]
    document = document_with("experience", [entry(fact_ids=[EXPERIENCE_ID, EXPERIENCE_ID_2])])

    outcome = validate_document(document, facts)

    codes = [f.code for f in outcome.findings]
    assert "ROLES_MERGED" in codes
    # The only entry failed, so the section goes too, with a finding saying so
    # rather than a heading left standing over nothing.
    assert outcome.document.sections == []
    assert "SECTION_OMITTED_EMPTY" in codes
    assert not outcome.passed


def test_a_project_moved_into_experience_is_refused() -> None:
    project = fact(
        PROJECT_ID,
        "project",
        {
            "name": "Analytical Engine",
            "role": None,
            "url": None,
            "start_month": None,
            "end_month": None,
            "bullets": [],
        },
    )
    document = document_with(
        "experience",
        [entry(title="Analytical Engine", organization=None, fact_ids=[PROJECT_ID])],
    )

    outcome = validate_document(document, [contact_fact(), project])

    assert "PROJECT_PRESENTED_AS_EMPLOYMENT" in [f.code for f in outcome.findings]
    assert outcome.document.sections == []


def test_a_skill_the_user_never_confirmed_is_refused() -> None:
    document = ResumeDocument(
        schema_version=1,
        language="en",
        contact=ResumeContact(
            full_name="Ada Lovelace",
            email=None,
            phone=None,
            location=None,
            links=[],
            fact_ids=[CONTACT_ID],
        ),
        sections=[
            ResumeSection(
                kind="skills",
                heading="Skills",
                entries=[
                    entry(title="Rust", organization=None, start_month=None, fact_ids=[SKILL_ID])
                ],
            )
        ],
    )

    outcome = validate_document(document, baseline_facts())

    assert "SKILL_NOT_CONFIRMED" in [f.code for f in outcome.findings]
    assert outcome.document.sections == []


def test_a_bullet_citing_an_unknown_fact_is_dropped() -> None:
    document = document_with(
        "experience",
        [entry(bullets=[ResumeBullet(text="Did the work", fact_ids=[UNKNOWN_ID])])],
    )

    outcome = validate_document(document, baseline_facts())

    assert "UNKNOWN_FACT_ID" in [f.code for f in outcome.findings]
    assert outcome.document.sections[0].entries[0].bullets == []


def test_an_entry_citing_an_unknown_fact_is_dropped_entirely() -> None:
    document = document_with("experience", [entry(fact_ids=[UNKNOWN_ID])])

    outcome = validate_document(document, baseline_facts())

    assert "UNKNOWN_FACT_ID" in [f.code for f in outcome.findings]
    assert outcome.document.sections == []


def test_a_credential_id_the_profile_lacks_is_dropped() -> None:
    certification = fact(
        CERT_ID,
        "certification",
        {
            "name": "Cloud Architect",
            "issuer": "Example Board",
            "credential_id": None,
            "issued_month": None,
            "expires_month": None,
        },
    )
    document = ResumeDocument(
        schema_version=1,
        language="en",
        contact=ResumeContact(
            full_name="Ada Lovelace",
            email=None,
            phone=None,
            location=None,
            links=[],
            fact_ids=[CONTACT_ID],
        ),
        sections=[
            ResumeSection(
                kind="certifications",
                heading="Certifications",
                entries=[
                    entry(
                        title="Cloud Architect",
                        organization="Example Board",
                        start_month=None,
                        current=False,
                        detail="AB-12345",
                        fact_ids=[CERT_ID],
                    )
                ],
            )
        ],
    )

    outcome = validate_document(document, [contact_fact(), certification])

    assert "CREDENTIAL_NOT_IN_FACTS" in [f.code for f in outcome.findings]
    assert outcome.document.sections[0].entries[0].detail is None


def test_validation_collects_every_cited_fact() -> None:
    facts = [contact_fact(), experience_fact(bullets=["Ran the pipeline"]), skill_fact()]
    document = build_document(facts, "en")

    outcome = validate_document(document, facts)

    assert set(outcome.fact_ids) == {CONTACT_ID, EXPERIENCE_ID, SKILL_ID}


def test_a_true_document_passes_untouched() -> None:
    facts = [contact_fact(), experience_fact(bullets=["Ran the ingestion pipeline"]), skill_fact()]
    document = build_document(facts, "en")

    outcome = validate_document(document, facts)

    assert outcome.passed
    assert outcome.findings == []
    assert outcome.document == document


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def test_docx_contains_the_document_text() -> None:
    facts = [contact_fact(), experience_fact(bullets=["Ran the ingestion pipeline"]), skill_fact()]
    document = build_document(facts, "en")

    data = render_docx(document)
    parsed = DocxDocument(io.BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)

    assert "Ada Lovelace" in text
    assert "Orbital Foods" in text
    assert "Ran the ingestion pipeline" in text
    assert "Python" in text


def test_docx_keeps_accents_intact() -> None:
    accented = fact(
        CONTACT_ID,
        "contact",
        {
            "full_name": "José Muñoz",
            "email": "jose@example.invalid",
            "phone": None,
            "city": "Asunción",
            "country": "Paraguay",
            "links": [],
        },
    )
    document = build_document([accented, skill_fact()], "es")

    parsed = DocxDocument(io.BytesIO(render_docx(document)))
    text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)

    assert "José Muñoz" in text
    assert "Asunción" in text
    assert "Habilidades" in text


def test_html_escapes_markup_rather_than_rendering_it() -> None:
    """Invariant 9: an imported CV is untrusted data."""
    hostile = fact(
        CONTACT_ID,
        "contact",
        {
            "full_name": "<script>alert(1)</script>",
            "email": "x@example.invalid",
            "phone": None,
            "city": None,
            "country": None,
            "links": [],
        },
    )
    markup = build_html(build_document([hostile, skill_fact()], "en"))

    assert "<script>alert(1)</script>" not in markup
    assert "&lt;script&gt;" in markup


def test_html_references_no_external_resource() -> None:
    document = build_document(baseline_facts(), "en")
    markup = build_html(document)

    for forbidden in ("<script", "<img", "<link", "http://", "@import", "url("):
        assert forbidden not in markup, forbidden


def test_html_never_shrinks_below_the_readable_minimum() -> None:
    markup = build_html(build_document(baseline_facts(), "en"))

    assert "font-size: 10.5pt" in markup
    for too_small in ("9pt", "8pt", "7pt", "6pt"):
        assert f"font-size: {too_small}" not in markup


@pytest.mark.parametrize("language", ["en", "es"])
def test_html_declares_the_document_language(language: str) -> None:
    markup = build_html(build_document(baseline_facts(), language))

    assert f'<html lang="{language}">' in markup
