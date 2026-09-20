"""Assert that each fixture really has the property its tests rely on.

Generating a fixture is not the same as generating a *correct* fixture: a
"scanned" PDF that happens to contain extractable text, or an "encrypted" one
that opens without a password, would make the acceptance tests pass for the
wrong reason. Run this after fixtures/generate.py.

    cd services/worker && uv run python ../../fixtures/verify.py
"""

from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

FIXTURES = Path(__file__).resolve().parent
CVS = FIXTURES / "cvs"
MODELS = FIXTURES / "model-responses"

failures: list[str] = []
checks = 0


def check(condition: bool, description: str) -> None:
    global checks
    checks += 1
    if condition:
        print(f"  ok    {description}")
    else:
        print(f"  FAIL  {description}")
        failures.append(description)


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def main() -> int:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    print("text-cv.pdf")
    text = pdf_text(CVS / "text-cv.pdf")
    check(len(PdfReader(CVS / "text-cv.pdf").pages) == 2, "has two pages")
    check("Ana Rivera" in text, "name is extractable")
    check("Northwind Logistics" in text, "employer is extractable")
    check("2022-03" in text, "start date is extractable")
    check("210ms" in text, "the numeric claim is extractable (AT10 needs it)")
    check(len(text) > 1000, "extraction is plausibly long")
    check(
        "Requires sponsorship for the United States." in text,
        "authorization statement is extractable (AT07)",
    )

    print("text-cv-es.pdf")
    es = pdf_text(CVS / "text-cv-es.pdf")
    check("Ingeniera de Backend Senior" in es, "Spanish title is extractable (AT27)")
    check("migracion" in es, "Spanish body text is extractable")

    print("prompt-injection-cv.pdf")
    injected = pdf_text(CVS / "prompt-injection-cv.pdf")
    check("IGNORE ALL PREVIOUS INSTRUCTIONS" in injected, "hostile text is present (AT09)")
    check("attacker.example.invalid" in injected, "exfiltration URL is present")
    check("Blair Okonkwo" in injected, "the genuine content is also present")

    print("scanned-cv.pdf")
    scanned = pdf_text(CVS / "scanned-cv.pdf")
    check(scanned.strip() == "", "yields no extractable text, so OCR_REQUIRED is correct (AT03)")
    check(len(PdfReader(CVS / "scanned-cv.pdf").pages) == 1, "is a readable single-page PDF")

    print("too-short.pdf")
    short = pdf_text(CVS / "too-short.pdf")
    check(0 < len(short.strip()) < 200, "extraction is non-empty but implausibly short (AT03)")

    print("encrypted-cv.pdf")
    reader = PdfReader(CVS / "encrypted-cv.pdf")
    check(reader.is_encrypted, "is reported as encrypted (AT03)")
    try:
        pages_readable = bool((reader.pages[0].extract_text() or "").strip())
    except Exception:
        pages_readable = False
    check(not pages_readable, "content is not readable without the password")
    check(reader.decrypt("fixture-password") != 0, "decrypts with the documented password")

    print("malformed.pdf")
    try:
        PdfReader(CVS / "malformed.pdf").pages[0]
        malformed_rejected = False
    except (PdfReadError, Exception):
        malformed_rejected = True
    check(malformed_rejected, "is rejected by the PDF reader (AT03)")

    print("mislabelled.docx")
    head = (CVS / "mislabelled.docx").read_bytes()[:5]
    check(head.startswith(b"%PDF"), "magic bytes say PDF despite the .docx extension")
    check(
        not zipfile.is_zipfile(CVS / "mislabelled.docx"),
        "is not a zip, so signature validation must reject it before parsing",
    )

    print("malformed.docx")
    check(
        not zipfile.is_zipfile(CVS / "malformed.docx"),
        "is a truncated archive and is rejected (AT03)",
    )

    print("text-cv.docx")
    from docx import Document

    document = Document(CVS / "text-cv.docx")
    paragraphs = "\n".join(p.text for p in document.paragraphs)
    check("Ana Rivera" in paragraphs, "name is in a paragraph")
    check(len(document.tables) == 1, "has exactly one table")
    table_text = "\n".join(
        cell.text for row in document.tables[0].rows for cell in row.cells
    )
    check("Northwind Logistics" in table_text, "employment history is inside the table")
    check(
        "Northwind Logistics" not in paragraphs,
        "the table content is NOT duplicated in paragraphs, so table parsing is required",
    )

    print("linkedin-export.txt")
    export = (CVS / "linkedin-export.txt").read_text(encoding="utf-8")
    check("Mar 2022 - Present" in export, "uses LinkedIn's month-name date format")
    check("Native or bilingual proficiency" in export, "carries LinkedIn language wording")

    print("ambiguous.txt")
    ambiguous = (CVS / "ambiguous.txt").read_text(encoding="utf-8")
    check("2021 to present" in ambiguous, "has a year-only date (AT03)")
    check("Previous Company" in ambiguous, "has an unusable employer name")
    check("wants to learn Go" in ambiguous, "has an aspirational skill that must not become a fact")

    print("model responses")
    text_cv = json.loads((MODELS / "parse_profile.text-cv.json").read_text(encoding="utf-8"))
    kinds = {fact["kind"] for fact in text_cv["draft_facts"]}
    check(
        kinds
        == {
            "contact",
            "experience",
            "education",
            "certification",
            "project",
            "authorization",
            "skill",
            "language",
        },
        "the canonical response covers every fact kind the CV states",
    )
    check(
        all(
            fact["value"]["user_declared_proficiency"] is None
            for fact in text_cv["draft_facts"]
            if fact["kind"] == "skill"
        ),
        "no skill carries an inferred proficiency",
    )
    check(
        all(
            bullet["evidence_reference"]
            for fact in text_cv["draft_facts"]
            if fact["kind"] in {"experience", "project"}
            for bullet in fact["value"]["bullets"]
        ),
        "every bullet carries an evidence reference",
    )
    us = [
        fact
        for fact in text_cv["draft_facts"]
        if fact["kind"] == "authorization" and fact["value"]["country"] == "US"
    ]
    check(len(us) == 1, "US authorization is proposed as its own fact")
    check(
        us[0]["value"]["authorized"] == "unknown",
        "US authorization stays unknown because the CV never states it (AT07)",
    )
    check(
        all(0.0 <= fact["confidence"] <= 1.0 for fact in text_cv["draft_facts"]),
        "confidences are in range",
    )

    injection = json.loads(
        (MODELS / "parse_profile.prompt-injection-cv.json").read_text(encoding="utf-8")
    )
    blob = json.dumps(injection)
    check("Stanford" not in blob, "the injected degree does not appear in the response (AT09)")
    check("PhD" not in blob, "the injected credential does not appear")
    check("attacker.example.invalid" not in blob, "the exfiltration URL does not appear")
    check(
        "15 years" not in blob,
        "the injected experience claim does not appear",
    )
    check(
        any(w["code"] == "PROMPT_INJECTION_TEXT_IGNORED" for w in injection["warnings"]),
        "the response warns that instruction-like text was ignored",
    )

    ambiguous_response = json.loads(
        (MODELS / "parse_profile.ambiguous.json").read_text(encoding="utf-8")
    )
    check(
        len(ambiguous_response["warnings"]) >= 3,
        "the ambiguous response reports its gaps rather than filling them",
    )
    check(
        not any(
            "Previous Company" in json.dumps(fact["value"])
            for fact in ambiguous_response["draft_facts"]
        ),
        "the unusable employer is not proposed as a fact",
    )
    check(
        not any(
            "Go" in json.dumps(fact["value"]) for fact in ambiguous_response["draft_facts"]
        ),
        "the aspirational skill is not proposed as a fact",
    )

    print("determinism")
    manifest = FIXTURES / "MANIFEST.sha256"
    # Excluded: the manifest cannot contain its own digest, the scripts are
    # inputs rather than outputs, and documentation is not a fixture --
    # editing the README must not look like a fixture changing.
    digests = {
        path.relative_to(FIXTURES).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(FIXTURES.rglob("*"))
        if path.is_file() and path.suffix not in {".py", ".md"} and path != manifest
    }
    lines = [f"{digest}  {name}\n" for name, digest in sorted(digests.items())]
    if manifest.exists():
        check(
            manifest.read_text(encoding="utf-8") == "".join(lines),
            "regenerating the fixtures reproduced identical bytes (AT11 relies on this)",
        )
    else:
        manifest.write_text("".join(lines), encoding="utf-8", newline="\n")
        print("  note  wrote MANIFEST.sha256 for the first time")

    print(f"\n{checks - len(failures)}/{checks} checks passed")
    if failures:
        print("\nFailures:")
        for description in failures:
            print(f"  - {description}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
